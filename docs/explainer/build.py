#!/usr/bin/env python3
"""Build the Prospecta explainer ("Prospecta: long-term memory, explained") from its one source.

ADAPTED from the house explainer builder, spire-venue docs/explainers/spire/build.py at commit
a69d1d79425583525cef616cf98afc3252b74453 (janus-infra/spire-venue; the builder files are unchanged since that
commit). Kept as it is: the slide kinds, the page frame, the theme (theme.css), the renderer and its layout lint
(render.mjs), the figure primitives (fig.py), the PPTX wrap and the acceptance gate. Changed for this repository
(see README.md "What differs from the Spire builder"): no screenshots (the `shot` kind is gone), no
spire-project privacy and day-word modules (the gate carries their word lists here), an agenda that takes any number
of parts, a diagram with an optional takeaway band, and paths. The Spire-original docstring follows.


  slides.py       every slide's words, figure, source line and speaker notes
  figures.py      the drawn figures (on fig.py's primitives), written to diagrams/*.svg by this build
  build.py        this: figures, HTML, then PDF and PNGs (render.mjs), PPTX, the gate
  render.mjs      the explainers' one renderer: headless Chromium, HTML -> PDF and one PNG
                  per slide, plus the layout lint; theme.css their one design

Usage:  python3 build.py [--out DIR] [--publish] [--copy-to DIR]
  --publish   copy the PDF and PPTX beside this file as prospecta-explainer-v1.{pdf,pptx}
  --copy-to   also copy them to DIR (for example ~/Documents/Prospecta)
Needs:  python-pptx, Pillow, tesseract, magick, pdftotext, pdfinfo, node, and Playwright
        (see README.md "Rebuild").

The build publishes nothing unless the acceptance gate passes:
  layout lint (render.mjs): no title wraps, nothing overlaps or leaves its slide,
    every slide numbered bottom right;
  every slide has a Source: line and speaker notes;
  page count: PDF pages = PPTX slides = slides in slides.py;
  privacy (credential and path words plus emails and hosts), house vocabulary,
    local paths and day words: on the OCR of every slide, the PDF text and
    the speaker notes.
"""
import argparse
import csv
import hashlib
import html
import importlib.util
import io
import re
import shutil
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import figures  # noqa: E402
from slides import SLIDES, TITLE  # noqa: E402

NAME = "prospecta-explainer-v1"
LABEL = {}  # status chips belong to the Spire deck; none here


EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
HOSTS = re.compile(r"\.postgres\.|azure\.com|tailscale|\b10\.\d+\.\d+\.\d+\b|192\.168\.|\b(roger|king|prince|chef)\b", re.I)
# Our own tools and roles, which the audience does not need.
VOCAB = re.compile(r"first\s*-?mate|crew\s*-?mate|second\s*-?mate|herdr|\bpi\b|tmux|no-?mistakes|\bwak(e|es|ing|en)\b|"
                   r"\bbrief\w*|captain", re.I)
LOCAL = re.compile(r"/home/|~/|/tmp/|treehouse|\.local/", re.I)
# The day words of spire-project tests/demo/deck/timeless.py: no relative or time-of-day words.
DAY = re.compile(r"\b(?:days?|today|tonight|tomorrow|yesterday|overnight|mornings?|evenings?|"
                 r"afternoons?|nights?|nightly|weekend)\b", re.I)
# The credential and path words of spire-project tests/demo/deck/masks.py that apply here.
SENSITIVE = re.compile(r"ghp_|\bsk-[A-Za-z0-9]{12,}|localhost|/home/|treehouse|firstmate|doppler|[0-9a-f]{8}-[0-9a-f]{4}-", re.I)

# ------------------------------------------------------------------ html
def esc(s):
    return html.escape(s, quote=False)


def foot(n, total, src):
    return (f'<div class="rule"></div><div class="foot"><span class="src"><b>Source:</b> {esc(src)}</span>'
            f'<span class="num">{n} / {total}</span></div>')


def header(s):
    h = f'<div class="kicker">{esc(s["kicker"])}</div><h1>{esc(s["title"])}</h1>'
    if s.get("sub"):
        h += f'<h2>{esc(s["sub"])}</h2>'
    return h


def chip(status):
    return f'<span class="chip {status}">{esc(LABEL[status])}</span>'


def render_slide(s, n, total):
    kind, f = s["kind"], foot(n, total, s["src"])
    if kind == "title":
        return (f'<section class="slide title"><div class="kicker">{esc(s["kicker"])}</div>'
                f'<h1 class="big">{esc(s["title"])}</h1><h2>{esc(s["sub"])}</h2>'
                f'<div class="presenter">{esc(s["presenter"])}</div>{f}</section>')
    if kind == "section":
        return (f'<section class="slide section"><div class="big">{esc(s["big"])}</div><div class="kicker">{esc(s["kicker"])}</div>'
                f'<h1>{esc(s["title"])}</h1><h2>{esc(s["sub"])}</h2>{f}</section>')
    if kind == "agenda":
        parts = "".join(f'<div><span>{esc(k)}</span><b>{esc(t)}</b><p>{esc(d)}</p></div>' for k, t, d in s["parts"])
        cols = s.get("cols", 2)
        legend = "".join(chip(c) for c in s.get("legend", []))
        legend = f'<div class="legend"><em>{esc(s["legend_text"])}</em>{legend}</div>' if s.get("legend_text") else ""
        return (f'<section class="slide">{header(s)}<div class="agenda n{cols}" style="grid-template-columns:repeat({cols},1fr)">{parts}</div>'
                f'{legend}{f}</section>')
    if kind == "statement":
        pts = "".join(f"<li>{p}</li>" for p in s["points"])
        return f'<section class="slide">{header(s)}<div class="body" style="top:{s.get("top", 270)}px"><ul class="points" style="font-size:35px">{pts}</ul></div>{f}</section>'
    if kind == "diagram":
        svg = (HERE / "diagrams" / s["svg"]).read_text()
        top = s.get("top", 290 if s.get("sub") else 245)
        lesson = f'<div class="lesson"><span>{esc(s["lesson"][0])}</span>{esc(s["lesson"][1])}</div>' if s.get("lesson") else ""
        bottom = 240 if lesson else 112
        return (f'<section class="slide">{header(s)}<div class="diagram" style="top:{top}px; bottom:{bottom}px">{svg}</div>'
                f'{lesson}{f}</section>')
    if kind == "cards":
        cols = s.get("cols", 3)
        cards = "".join(f'<div class="card" style="border-top:10px solid {c}"><h3 style="color:{c}">{esc(h)}</h3><p>{p}</p></div>'
                        for h, p, c in s["cards"])
        lesson = f'<div class="lesson"><span>{esc(s["lesson"][0])}</span>{esc(s["lesson"][1])}</div>' if s.get("lesson") else ""
        bottom = 250 if lesson else 112
        return (f'<section class="slide">{header(s)}<div class="body cards" style="top:{s.get("top", 260)}px; bottom:{bottom}px; '
                f'grid-template-columns:repeat({cols},1fr); align-content:stretch">{cards}</div>{lesson}{f}</section>')
    if kind == "twocol":
        def col(c):
            h, colr, pts = c
            li = "".join(f"<li>{p}</li>" for p in pts)
            return f'<div class="card" style="border-top:10px solid {colr}"><h3 style="color:{colr}">{esc(h)}</h3><ul class="points" style="font-size:32px">{li}</ul></div>'
        return (f'<section class="slide">{header(s)}<div class="body cards" style="top:{s.get("top", 260)}px; grid-template-columns:1fr 1fr; align-content:stretch">'
                f'{col(s["left"])}{col(s["right"])}</div>{f}</section>')
    if kind == "map":
        rows = "".join(f'<tr><td class="ours">{esc(a)}</td><td class="arrow">→</td><td class="spire">{esc(b)}</td>'
                       f'<td><span class="mchip {c}">{esc(t)}</span></td></tr>' for a, b, c, t in s["rows"])
        return (f'<section class="slide">{header(s)}<div class="maptable"><table><thead><tr><th>{esc(s["head"][0])}</th><th></th>'
                f'<th>{esc(s["head"][1])}</th><th>{esc(s["head"][2])}</th></tr></thead><tbody>{rows}</tbody></table></div>{f}</section>')
    if kind == "table":
        rows = "".join("<tr>" + "".join(f"<td>{esc(c)}</td>" for c in r) + "</tr>" for r in s["rows"])
        head = "".join(f"<th>{esc(h)}</th>" for h in s["head"])
        return (f'<section class="slide">{header(s)}<div class="tbl"><table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table></div>'
                f'<div class="tfoot">{esc(s["foot"])}</div>{f}</section>')
    if kind == "status":
        rows = "".join(f'<tr><td class="what">{esc(a)}</td><td>{esc(b)}</td><td class="st">{chip(c)}</td></tr>' for a, b, c in s["rows"])
        return f'<section class="slide">{header(s)}<div class="tbl status"><table><tbody>{rows}</tbody></table></div>{f}</section>'
    if kind == "closing":
        ev = "".join(f'<div class="ev"><div class="evn">{esc(big)}</div><div class="evt">{esc(t)}</div></div>' for big, t in s["evidence"])
        return (f'<section class="slide title closing"><div class="kicker">{esc(s["kicker"])}</div><h1>{esc(s["title"])}</h1>'
                f'<h2>{esc(s["sub"])}</h2><div class="evrow">{ev}</div>{f}</section>')
    raise ValueError(f"unknown kind {kind}")


EXTRA_CSS = """
.slide.title h1.big { font-size:200px; margin-top:150px; margin-bottom:40px; letter-spacing:-0.02em; }
.slide.title .presenter { position:absolute; left:160px; bottom:150px; font-size:34px; color:#c9cbe0; }
.slide.section h1 { margin-top:140px; }
.card .points li { margin-bottom:14px; }
.cards .card p { font-size:29px; }
.cards .card h3 { font-size:32px; }
.card .points li::before { top:13px; }
.lesson { position:absolute; left:120px; right:120px; bottom:120px; height:100px; background:var(--ink); color:#fff; border-radius:16px; font-size:27px; display:flex; align-items:center; padding:0 36px; gap:28px; }
.lesson span { font-size:18px; letter-spacing:.14em; text-transform:uppercase; color:#b3a9ff; font-weight:700; white-space:nowrap; }
.agenda { position:absolute; left:120px; right:120px; top:280px; display:grid; grid-template-columns:1fr 1fr; gap:36px; }
.agenda div { background:#fff; border-radius:18px; padding:36px 40px; box-shadow:0 0 0 1px var(--line); min-height:380px; overflow:hidden; }
.agenda span { display:block; font-size:120px; font-weight:800; color:var(--accent); line-height:1; opacity:.3; }
.agenda b { display:block; font-size:44px; margin:10px 0 18px; }
.agenda p { font-size:29px; line-height:1.4; color:var(--ink2); margin:0; }
.agenda.n3 { gap:22px; top:270px; }
.agenda.n3 div { min-height:0; padding:18px 28px; }
.agenda.n3 span { font-size:54px; }
.agenda.n3 b { font-size:34px; margin:2px 0 6px; }
.agenda.n3 p { font-size:24px; line-height:1.3; }
.legend { position:absolute; left:120px; right:120px; top:740px; display:flex; gap:18px; align-items:center; font-size:24px; color:var(--ink2); }
.legend em { font-style:normal; margin-right:8px; }
.chip { display:inline-block; padding:7px 18px; border-radius:999px; font-size:21px; font-weight:700; color:#fff; white-space:nowrap; }
.chip.built { background:var(--hit); } .chip.owed, .chip.unproven { background:var(--detour); } .chip.spec { background:var(--missed); } .chip.code { background:var(--ink2); }
.maptable { position:absolute; left:120px; right:120px; top:250px; bottom:100px; overflow:hidden; }
.maptable table { width:100%; border-collapse:separate; border-spacing:0 8px; font-size:25px; }
.maptable th { text-align:left; font-size:17px; letter-spacing:.12em; text-transform:uppercase; color:var(--muted); font-weight:700; padding:0 18px; }
.maptable td { background:#fff; padding:12px 18px; line-height:1.25; box-shadow:0 1px 0 var(--line); }
.maptable td.ours { width:32%; border-radius:12px 0 0 12px; font-weight:600; }
.maptable td.arrow { width:3%; color:var(--accent); font-weight:800; padding:0; text-align:center; }
.maptable td.spire { width:44%; color:var(--ink2); }
.maptable td:last-child { border-radius:0 12px 12px 0; width:21%; }
.mchip { display:inline-block; padding:5px 14px; border-radius:999px; font-size:19px; font-weight:700; color:#fff; white-space:nowrap; }
.mchip.same { background:var(--detour); } .mchip.ahead { background:var(--hit); } .mchip.differs { background:var(--ink2); }
.tbl { position:absolute; left:120px; right:120px; top:250px; overflow:hidden; }
.tbl table { width:100%; border-collapse:separate; border-spacing:0 9px; font-size:27px; }
.tbl th { text-align:left; font-size:18px; letter-spacing:.12em; text-transform:uppercase; color:var(--muted); font-weight:700; padding:0 22px; }
.tbl td { background:#fff; padding:18px 22px; line-height:1.3; box-shadow:0 1px 0 var(--line); vertical-align:middle; }
.tbl td:first-child { border-radius:12px 0 0 12px; font-weight:700; width:34%; }
.tbl td:last-child { border-radius:0 12px 12px 0; }
.tbl.status table { font-size:23px; border-spacing:0 8px; }
.tbl.status td { padding:12px 20px; }
.tbl.status td.what { width:36%; }
.tbl.status td.st { width:17%; text-align:right; }
.tbl.status .chip { font-size:18px; padding:6px 14px; }
.tfoot { position:absolute; left:120px; right:120px; bottom:118px; font-size:26px; line-height:1.4; color:var(--ink); background:var(--accent-soft); border-radius:14px; padding:20px 28px; }
.shotbox { position:absolute; left:120px; top:250px; width:1088px; height:680px; border-radius:12px; overflow:hidden;
  box-shadow:0 18px 50px rgba(20,24,40,.16), 0 0 0 1px rgba(20,24,40,.12); background:#fff; }
.shotbox img { width:1088px; height:680px; display:block; }
.caption2 { position:absolute; left:120px; top:942px; width:1088px; font-size:17px; color:var(--muted); line-height:1.3; }
.side { position:absolute; left:1250px; right:110px; top:250px; bottom:112px; overflow:hidden; }
.side .said { background:var(--card); border-left:8px solid var(--ink); border-radius:10px; padding:20px 24px; font-size:27px; line-height:1.35; box-shadow:0 2px 0 rgba(0,0,0,.04); }
.side .said .by { display:block; font-size:17px; text-transform:uppercase; letter-spacing:.1em; color:var(--muted); margin-bottom:8px; }
.side .then { margin-top:24px; font-size:25px; line-height:1.42; color:var(--ink); }
.side .then b { color:var(--accent); }
.closing h1 { margin-top:40px; }
.evrow { position:absolute; left:160px; right:160px; top:620px; display:grid; grid-template-columns:1fr 1fr 1fr; gap:36px; }
.ev { background:#273049; border-radius:18px; padding:30px 32px; }
.evn { font-size:62px; font-weight:800; color:#b3a9ff; line-height:1; margin-bottom:18px; }
.evt { font-size:25px; line-height:1.4; color:#e6e7f2; }
"""


def build_html(out):
    total = len(SLIDES)
    body = "\n".join(render_slide(s, i + 1, total) for i, s in enumerate(SLIDES))
    css = (HERE / "theme.css").read_text() + EXTRA_CSS
    (out / "deck.html").write_text(f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{esc(TITLE)}</title>'
                                   f'<style>{css}</style></head><body>{body}</body></html>')


# ------------------------------------------------------------------ ocr
def ocr_tsv(png, work):
    digest = hashlib.sha1(png.read_bytes()).hexdigest()[:16]
    tsv = work / f"{png.stem}-{digest}.tsv"
    if not tsv.exists():
        big = work / (png.stem + "-2x.png")
        subprocess.run(["magick", str(png), "-resize", "200%", str(big)], check=True)
        subprocess.run(["nice", "tesseract", str(big), str(tsv.with_suffix("")), "--psm", "11", "-c", "tessedit_create_tsv=1", "quiet"],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        big.unlink()
    return tsv.read_text()


def sensitive_finds(tsv_text):
    return [(r.get("text") or "") for r in csv.DictReader(io.StringIO(tsv_text), delimiter="\t", quoting=csv.QUOTE_NONE)
            if (r.get("text") or "").strip() and SENSITIVE.search(r.get("text") or "")]


# ------------------------------------------------------------------ pptx
def notes_text(s):
    return s["notes"] + f"\n\nSource: {s['src']}"


def build_pptx(out, dest):
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    for i, s in enumerate(SLIDES):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        sl.shapes.add_picture(str(out / "png" / f"{i + 1:02d}.png"), 0, 0, prs.slide_width, prs.slide_height)
        sl.notes_slide.notes_text_frame.text = notes_text(s)
    prs.core_properties.title = f"{TITLE}: long-term memory, explained"
    prs.core_properties.author = "Prospecta"
    prs.save(dest)


# ------------------------------------------------------------------ the gate
def bad_word(t):
    for name, rx in (("private", SENSITIVE), ("email", EMAIL), ("host", HOSTS), ("vocabulary", VOCAB),
                     ("local path", LOCAL), ("day word", DAY)):
        m = rx.search(t)
        if m:
            return f"{name}: {m.group(0)!r}"
    return None


def gate(out, pptx_path):
    """Every failure of the acceptance gate, as (where, what)."""
    bad = []
    for i, s in enumerate(SLIDES, 1):
        if not s.get("src", "").strip():
            bad.append((f"slide {i}", "no Source: line"))
        if not s.get("notes", "").strip():
            bad.append((f"slide {i}", "no speaker notes"))
    pages = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(out / "deck.pdf")], capture_output=True, text=True,
                                                             check=True).stdout).group(1))
    from pptx import Presentation
    prs = Presentation(pptx_path)
    if not (pages == len(prs.slides) == len(SLIDES)):
        bad.append(("page count", f"PDF {pages}, PPTX {len(prs.slides)}, slides.py {len(SLIDES)}"))
    work = out / "ocr"
    work.mkdir(exist_ok=True)
    for png in sorted((out / "png").glob("*.png")):
        tsv = ocr_tsv(png, work)
        for name in sensitive_finds(tsv):
            bad.append((f"ocr {png.name}", f"private: {name!r}"))
        for r in csv.DictReader(io.StringIO(tsv), delimiter="\t", quoting=csv.QUOTE_NONE):
            t = (r.get("text") or "").strip()
            if not t:
                continue
            rules = (EMAIL, HOSTS, VOCAB, LOCAL, DAY)
            if any(rx.search(t) for rx in rules):
                bad.append((f"ocr {png.name}", t))
    pdf_text = subprocess.run(["pdftotext", str(out / "deck.pdf"), "-"], capture_output=True, text=True, check=True).stdout
    for line in pdf_text.splitlines():
        why = bad_word(line)
        if why:
            bad.append(("pdf text", f"{why} in {line.strip()!r}"))
    for i, sl in enumerate(prs.slides, 1):
        notes = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
        if not notes.strip():
            bad.append((f"pptx slide {i}", "no speaker notes"))
        for line in notes.splitlines():
            why = bad_word(line)
            if why:
                bad.append((f"pptx notes {i}", f"{why} in {line.strip()[:90]!r}"))
    return bad, pages


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(HERE / "out"))
    ap.add_argument("--publish", action="store_true", help=f"copy the PDF and PPTX beside this file as {NAME}.{{pdf,pptx}}")
    ap.add_argument("--copy-to", default=None, help=f"also copy them to this directory as {NAME}.{{pdf,pptx}}")
    a = ap.parse_args()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    for old in (out / "png").glob("*.png"):
        old.unlink()
    figures.write_all(HERE / "diagrams")
    build_html(out)
    subprocess.run(["nice", "node", str(HERE / "render.mjs"), str(out), "--require-source"], check=True)
    build_pptx(out, out / "deck.pptx")
    bad, pages = gate(out, out / "deck.pptx")
    if bad:
        for where, word in bad:
            print(f"GATE: {where}: {word}", file=sys.stderr)
        sys.exit("acceptance gate failed: nothing published")
    print(f"gate: clean on {len(SLIDES)} slides ({pages} PDF pages = PPTX slides; layout, sources, notes; "
          "privacy, vocabulary, local paths and day words on OCR, PDF text and notes)")
    dests = ([HERE] if a.publish else []) + ([Path(a.copy_to).expanduser()] if a.copy_to else [])
    for dest in dests:
        shutil.copy(out / "deck.pdf", dest / f"{NAME}.pdf")
        shutil.copy(out / "deck.pptx", dest / f"{NAME}.pptx")
        print(f"published {NAME}.pdf and .pptx to {dest}")


if __name__ == "__main__":
    main()
